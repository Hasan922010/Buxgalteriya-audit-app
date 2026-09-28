import io
import uuid
from datetime import date
from decimal import Decimal
import openpyxl
import pytest

from app.schemas.report import MaterialReport, MaterialReportItem, MaterialReportMxikGroup
from app.services.export_engine import ExportEngine
from app.modules.reports.excel_export import ExcelExportEngine


def create_sample_material_report():
    org_id = uuid.uuid4()
    item1 = MaterialReportItem(
        item_id=uuid.uuid4(),
        item_name="Sement M-400 (50kg)",
        ikpu_code="02501001001000000",
        unit="qop",
        initial_qty=Decimal("100.000"),
        initial_sum=Decimal("6000000.00"),
        inflow_qty=Decimal("200.000"),
        inflow_sum=Decimal("12000000.00"),
        outflow_qty=Decimal("150.000"),
        outflow_sum=Decimal("9000000.00"),
        avg_price=Decimal("60000.00"),
        final_qty=Decimal("150.000"),
        final_sum=Decimal("9000000.00")
    )
    item2 = MaterialReportItem(
        item_id=uuid.uuid4(),
        item_name="Sement M-500 (50kg)",
        ikpu_code="02501001001000000",
        unit="qop",
        initial_qty=Decimal("50.000"),
        initial_sum=Decimal("3500000.00"),
        inflow_qty=Decimal("100.000"),
        inflow_sum=Decimal("7000000.00"),
        outflow_qty=Decimal("80.000"),
        outflow_sum=Decimal("5600000.00"),
        avg_price=Decimal("70000.00"),
        final_qty=Decimal("70.000"),
        final_sum=Decimal("4900000.00")
    )

    mxik1 = MaterialReportMxikGroup(
        ikpu_code="02501001001000000",
        ikpu_name="Sement va gips mahsulotlari",
        items_count=2,
        unit="qop",
        initial_qty=Decimal("150.000"),
        initial_sum=Decimal("9500000.00"),
        inflow_qty=Decimal("300.000"),
        inflow_sum=Decimal("19000000.00"),
        outflow_qty=Decimal("230.000"),
        outflow_sum=Decimal("14600000.00"),
        final_qty=Decimal("220.000"),
        final_sum=Decimal("13900000.00")
    )

    return MaterialReport(
        organization_id=org_id,
        organization_name="TEST STROY MCHJ",
        from_date=date(2025, 1, 1),
        to_date=date(2025, 1, 31),
        items=[item1, item2],
        mxik_groups=[mxik1],
        total_initial_sum=Decimal("9500000.00"),
        total_inflow_sum=Decimal("19000000.00"),
        total_outflow_sum=Decimal("14600000.00"),
        total_final_sum=Decimal("13900000.00")
    )


def test_export_engine_material_excel_has_mxik_tovarlar_soni():
    report = create_sample_material_report()
    stream = ExportEngine.export_material_excel(report)
    wb = openpyxl.load_workbook(stream)

    assert "Moddiy_Hisobot" in wb.sheetnames
    assert "MXIK_Jamlama" in wb.sheetnames

    ws_mxik = wb["MXIK_Jamlama"]

    # Verify hierarchical headers
    # Row 4 headers
    assert "MXIK (IKPU) Kodi" in str(ws_mxik["B4"].value)
    assert "Tovarlar turlari soni" in str(ws_mxik["E4"].value)
    assert "Boshlang'ich qoldiq" in str(ws_mxik["F4"].value)
    assert "Davr kirimi" in str(ws_mxik["H4"].value)
    assert "Davr chiqimi" in str(ws_mxik["J4"].value)
    assert "Yakuniy qoldiq" in str(ws_mxik["L4"].value)

    # Row 5 sub-headers have "Tovarlar soni" for quantities
    assert "Tovarlar soni" in str(ws_mxik["F5"].value)
    assert "Summasi (so'm)" in str(ws_mxik["G5"].value)
    assert "Tovarlar soni" in str(ws_mxik["H5"].value)
    assert "Summasi (so'm)" in str(ws_mxik["I5"].value)
    assert "Tovarlar soni" in str(ws_mxik["J5"].value)
    assert "Summasi (so'm)" in str(ws_mxik["K5"].value)
    assert "Tovarlar soni" in str(ws_mxik["L5"].value)
    assert "Summasi (so'm)" in str(ws_mxik["M5"].value)

    # Row 6: First MXIK data row
    assert ws_mxik["B6"].value == "02501001001000000"
    assert ws_mxik["C6"].value == "Sement va gips mahsulotlari"
    assert ws_mxik["D6"].value == "qop"
    assert ws_mxik["E6"].value == 2       # items_count (Tovarlar turlari soni)
    assert ws_mxik["F6"].value == 150.0   # initial_qty (Boshlang'ich tovarlar soni)
    assert ws_mxik["G6"].value == 9500000.0 # initial_sum
    assert ws_mxik["H6"].value == 300.0   # inflow_qty (Kirim tovarlar soni)
    assert ws_mxik["I6"].value == 19000000.0 # inflow_sum
    assert ws_mxik["J6"].value == 230.0   # outflow_qty (Chiqim tovarlar soni)
    assert ws_mxik["K6"].value == 14600000.0 # outflow_sum
    assert ws_mxik["L6"].value == 220.0   # final_qty (Yakuniy qoldiq tovarlar soni)
    assert ws_mxik["M6"].value == 13900000.0 # final_sum

    # Row 7: Totals row
    assert "JAMI:" in str(ws_mxik["A7"].value)
    assert ws_mxik["E7"].value == 2       # Total items count
    assert ws_mxik["F7"].value == 150.0   # Total initial qty
    assert ws_mxik["H7"].value == 300.0   # Total inflow qty
    assert ws_mxik["J7"].value == 230.0   # Total outflow qty
    assert ws_mxik["L7"].value == 220.0   # Total final qty


def test_excel_export_engine_material_report_has_mxik_tovarlar_soni():
    report = create_sample_material_report()
    stream = ExcelExportEngine.export_material_report(report)
    wb = openpyxl.load_workbook(stream)

    assert "Moddiy_Hisobot" in wb.sheetnames
    assert "MXIK_Jamlama" in wb.sheetnames

    ws_mxik = wb["MXIK_Jamlama"]
    assert "Tovarlar turlari soni" in str(ws_mxik["E4"].value)
    assert "Tovarlar soni" in str(ws_mxik["F5"].value)
    assert ws_mxik["E6"].value == 2
    assert ws_mxik["F6"].value == 150.0
    assert ws_mxik["H6"].value == 300.0
    assert ws_mxik["J6"].value == 230.0
    assert ws_mxik["L6"].value == 220.0
