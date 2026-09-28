import pytest
import io
from decimal import Decimal
import openpyxl
from httpx import AsyncClient, ASGITransport

from app.main import app
from app.modules.ocr.tax_audit_schemas import TaxAuditDocument, TaxAuditItemRow, TaxPenaltySummary
from app.modules.ocr.tax_audit_extractor import TaxAuditExtractor
from app.modules.ocr.tax_excel_generator import TaxAuditExcelGenerator


def test_clean_decimal_handles_messy_numbers():
    assert TaxAuditExtractor.clean_decimal("20 025 600,00") == Decimal("20025600.00")
    assert TaxAuditExtractor.clean_decimal(" 1 500,50 ") == Decimal("1500.50")
    assert TaxAuditExtractor.clean_decimal("0,00") == Decimal("0.0")
    assert TaxAuditExtractor.clean_decimal(None) == Decimal("0.0")


def test_validate_row_math_discrepancy_and_penalties():
    # Item where sold (200) > available (50 + 50 = 100), discrepancy = 100
    row = TaxAuditItemRow(
        item_no=1,
        item_name="Sement M-400",
        opening_qty=Decimal("50"),
        opening_amount=Decimal("3000000"),
        inflow_qty=Decimal("50"),
        inflow_amount=Decimal("3000000"),
        sold_qty=Decimal("200"),
        avg_price=Decimal("70000"),
        sold_amount=Decimal("14000000"),
        closing_qty=Decimal("0"),
        closing_amount=Decimal("0"),
        diff_qty=Decimal("0"),
        diff_amount=Decimal("0"),
    )

    validated = TaxAuditExtractor.validate_row_math(row)
    assert validated.diff_qty == Decimal("100")
    assert validated.diff_amount == Decimal("7000000")  # 100 * 70,000
    assert validated.math_verified is True

    # Calculate penalties for this row
    summary = TaxAuditExtractor.calculate_penalties([validated])
    assert summary.total_discrepancy_amount == Decimal("7000000")
    # VAT (12% of 7,000,000 included): 7,000,000 * 0.12 / 1.12 = 750,000.00
    assert summary.vat_amount == Decimal("750000.00")
    # Net tax base: 7,000,000 - 750,000 = 6,250,000.00
    assert summary.net_tax_base == Decimal("6250000.00")
    # Profit tax (15%): 6,250,000 * 0.15 = 937,500.00
    assert summary.profit_tax_addition == Decimal("937500.00")
    # Financial penalty (20%): 7,000,000 * 0.20 = 1,400,000.00
    assert summary.financial_penalty == Decimal("1400000.00")
    # Total budget liability: 750,000 + 937,500 + 1,400,000 = 3,087,500.00
    assert summary.total_budget_liability == Decimal("3087500.00")


def test_tax_excel_generator_structure_and_formulas():
    row1 = TaxAuditItemRow(
        item_no=1,
        item_name="Sement M-400 (50kg)",
        opening_qty=Decimal("100"),
        opening_amount=Decimal("6000000"),
        inflow_qty=Decimal("300"),
        inflow_amount=Decimal("18000000"),
        sold_qty=Decimal("550"),
        avg_price=Decimal("65000"),
        sold_amount=Decimal("35750000"),
        closing_qty=Decimal("0"),
        closing_amount=Decimal("0"),
        diff_qty=Decimal("150"),
        diff_amount=Decimal("9750000"),
    )
    summary = TaxAuditExtractor.calculate_penalties([row1])

    doc = TaxAuditDocument(
        company_name="OOO GRAND STROY",
        company_inn="301234567",
        audit_year=2023,
        title="Киримсиз товарлар таҳлили",
        items=[row1],
        summary=summary
    )

    excel_bytes = TaxAuditExcelGenerator.generate_workbook(doc)
    assert len(excel_bytes) > 2000
    assert excel_bytes[:4] == b"PK\x03\x04"

    # Open with openpyxl to inspect cells
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Киримсиз товарлар таҳлили"

    # Title check
    assert "OOO GRAND STROY" in str(ws["A1"].value)
    assert "301234567" in str(ws["A1"].value)

    # Row 5 (first data item)
    assert ws["B5"].value == "Sement M-400 (50kg)"
    assert ws["L5"].value == 150.0  # diff_qty
    assert ws["M5"].value == 9750000.0  # diff_amount
    # Check red alert fill on discrepancy
    assert ws["L5"].fill.start_color.rgb == "00FFF0F0"

    # Totals row (row 6)
    assert "=SUM(L5:L5)" == ws["L6"].value
    assert "=SUM(M5:M5)" == ws["M6"].value


@pytest.mark.asyncio
async def test_tax_audit_api_endpoints():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # Test 1: parse endpoint with dummy pdf/text
        files = {"file": ("tax_audit.pdf", b"%PDF-1.4 dummy test content", "application/pdf")}
        res = await ac.post("/api/v1/ocr/tax-audit/parse", files=files)
        assert res.status_code == 200
        data = res.json()
        assert "company_name" in data
        assert "items" in data
        assert len(data["items"]) >= 1
        assert "summary" in data
        assert data["summary"]["total_discrepancy_amount"] is not None

        # Test 2: export-excel endpoint with the parsed document
        res_excel = await ac.post("/api/v1/ocr/tax-audit/export-excel", json=data)
        assert res_excel.status_code == 200
        assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert len(res_excel.content) > 1000
        assert res_excel.content[:4] == b"PK\x03\x04"
