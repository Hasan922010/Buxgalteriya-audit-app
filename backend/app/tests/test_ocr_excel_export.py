import io
import tempfile
from datetime import date
from decimal import Decimal
import pytest
import openpyxl
from httpx import AsyncClient, ASGITransport

from app.main import app
from app.modules.ocr.ocr_extractor import ExtractedDocument, ExtractedLineItem
from app.modules.ocr.excel_exporter import OCRExcelExporter
from app.modules.documents.parsers.didox_parser import DidoxParser

@pytest.fixture
def sample_ocr_document() -> ExtractedDocument:
    return ExtractedDocument(
        doc_number="EHF-2025-088",
        doc_date=date(2025, 5, 15),
        doc_type="EHF",
        supplier_name="Sanoat Mega Savdo MCHJ",
        supplier_inn="305123456",
        buyer_name="Toshkent Qurilish Invest OK",
        buyer_inn="308987654",
        contract_number="42-A",
        contract_date=date(2025, 5, 10),
        line_items=[
            ExtractedLineItem(
                item_name="Sement M-500 qoplarda",
                ikpu_code="02310001001000000",
                unit="tonna",
                quantity=Decimal("15.5"),
                price=Decimal("850000.00"),
                vat_rate=Decimal("12.0"),
                vat_amount=Decimal("1581000.00"),
                total_amount=Decimal("14756000.00"),
                confidence=0.98
            ),
            ExtractedLineItem(
                item_name="Armatura d12 A500C",
                ikpu_code="02410002002000000",
                unit="tonna",
                quantity=Decimal("5.0"),
                price=Decimal("9200000.00"),
                vat_rate=Decimal("12.0"),
                vat_amount=Decimal("5520000.00"),
                total_amount=Decimal("51520000.00"),
                confidence=0.95
            ),
            ExtractedLineItem(
                item_name="Gisht pishiq standart",
                ikpu_code="02320003001000000",
                unit="dona",
                quantity=Decimal("5000.0"),
                price=Decimal("1800.00"),
                vat_rate=Decimal("12.0"),
                vat_amount=Decimal("1080000.00"),
                total_amount=Decimal("10080000.00"),
                confidence=0.92
            ),
        ]
    )

def test_ocr_excel_exporter_structure(sample_ocr_document):
    """
    Verifies that the generated Excel workbook is visually styled like an authentic
    Uzbekistan electronic invoice (Didox EHF) with proper title, requisites, table, and signatures.
    """
    buf = OCRExcelExporter.export_document_to_excel(sample_ocr_document)
    assert buf is not None
    assert len(buf.getvalue()) > 2000

    wb = openpyxl.load_workbook(buf)
    assert "Hisobvaraq_Faktura" in wb.sheetnames
    ws = wb["Hisobvaraq_Faktura"]

    # Check Title
    assert "EHF-2025-088" in str(ws["A1"].value)
    assert "ELEKTRON HISOBVARAQ-FAKTURA" in str(ws["A1"].value)

    # Check Supplier & Buyer cards
    assert "Sanoat Mega Savdo MCHJ" in str(ws["A4"].value)
    assert "305123456" in str(ws["A5"].value)
    assert "Toshkent Qurilish Invest OK" in str(ws["H4"].value)
    assert "308987654" in str(ws["H5"].value)

    # Check Table Headers at Row 7
    header_vals = [cell.value for cell in ws[7]]
    assert "№" in header_vals
    assert "Tovarlar (xizmatlar) nomi" in header_vals
    assert "IKPU / MXIK kodi" in header_vals
    assert "Jami qiymat" in header_vals

    # Check Items Data
    assert ws["H8"].value == "Sement M-500 qoplarda"
    assert ws["H9"].value == "Armatura d12 A500C"
    assert ws["H10"].value == "Gisht pishiq standart"

    # Check Total row exists
    total_cell = ws.cell(row=11, column=1)
    assert "JAMI" in str(total_cell.value)
    # Total sum is 14756000 + 51520000 + 10080000 = 76356000
    grand_total_cell = ws.cell(row=11, column=16)
    assert float(grand_total_cell.value) == 76356000.0

@pytest.mark.asyncio
async def test_ocr_export_excel_api_endpoint(sample_ocr_document):
    """
    Tests POST /api/v1/ocr/export-excel endpoint.
    """
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        payload = sample_ocr_document.model_dump(mode="json")
        resp = await client.post("/api/v1/ocr/export-excel", json=payload)
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert "attachment" in resp.headers["content-disposition"]
        assert "faktura_EHF-2025-088_aslidek.xlsx" in resp.headers["content-disposition"]

        # Verify the returned bytes can be loaded as openpyxl Workbook
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "Hisobvaraq_Faktura"
        assert "Sement M-500 qoplarda" in [cell.value for cell in ws["H"]]

def test_ocr_exported_excel_roundtrip_with_didox_parser(sample_ocr_document):
    """
    CRITICAL USER REQUIREMENT VERIFICATION:
    Verifies that the generated authentic .xlsx can be ingested back directly by DidoxParser
    without any modification ("aslidek qilib saqlab bersin keyin o'zim tekshirib kiritaman").
    """
    buf = OCRExcelExporter.export_document_to_excel(sample_ocr_document)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tf.write(buf.getvalue())
        tf_path = tf.name

    parser = DidoxParser()
    records = parser.parse_file(tf_path)

    # Exactly 3 line items parsed
    assert len(records) == 3

    # Verify Item 1
    r1 = records[0]
    assert r1.item_name == "Sement M-500 qoplarda"
    assert r1.ikpu_code == "02310001001000000"
    assert r1.quantity == Decimal("15.5")
    assert r1.price == Decimal("850000.00")
    assert r1.total_amount == Decimal("14756000.00")
    assert r1.supplier_inn == "305123456"
    assert r1.buyer_inn == "308987654"

    # Verify Item 2
    r2 = records[1]
    assert r2.item_name == "Armatura d12 A500C"
    assert r2.quantity == Decimal("5.0")
    assert r2.total_amount == Decimal("51520000.00")

    # Verify Item 3
    r3 = records[2]
    assert r3.item_name == "Gisht pishiq standart"
    assert r3.quantity == Decimal("5000.0")
    assert r3.total_amount == Decimal("10080000.00")
