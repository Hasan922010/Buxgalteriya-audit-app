import io
from decimal import Decimal
from typing import Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from app.schemas.report import TrialBalanceReport, MaterialReport, AktSverkaReport

class ExportEngine:
    """
    Generates styled Excel and PDF exports for accounting reports.
    """

    @staticmethod
    def export_oborotka_excel(report: TrialBalanceReport) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "OSV_Oborotka"

        # Fonts & Fills
        title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        # Title rows
        ws.append([f"Aylanma Qoldiq Vedomosti (OSV) - {report.organization_name}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
        ws.cell(row=2, column=1).font = Font(size=10, italic=True)
        ws.append([])

        # Table Headers
        headers_row1 = ["Schot", "Schot nomi", "Turi", "Boshlang'ich qoldiq", "", "Davr oboroti", "", "Oxirgi qoldiq", ""]
        headers_row2 = ["", "", "", "Debet", "Kredit", "Debet", "Kredit", "Debet", "Kredit"]
        
        ws.append(headers_row1)
        ws.append(headers_row2)

        start_row = 4
        # Merges for headers
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row+1, end_column=1)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row+1, end_column=2)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row+1, end_column=3)
        ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=5)
        ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row, end_column=7)
        ws.merge_cells(start_row=start_row, start_column=8, end_row=start_row, end_column=9)

        for r in range(start_row, start_row + 2):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        # Data rows
        curr_row = start_row + 2
        for it in report.items:
            row_data = [
                it.account_code,
                it.account_name,
                it.account_type,
                float(it.initial_debit),
                float(it.initial_credit),
                float(it.turnover_debit),
                float(it.turnover_credit),
                float(it.final_debit),
                float(it.final_credit),
            ]
            ws.append(row_data)
            for col_idx in range(1, 10):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in [1, 3]:
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx >= 4:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
            curr_row += 1

        # Totals Row
        totals = [
            "JAMI",
            "",
            "",
            float(report.total_initial_debit),
            float(report.total_initial_credit),
            float(report.total_turnover_debit),
            float(report.total_turnover_credit),
            float(report.total_final_debit),
            float(report.total_final_credit),
        ]
        ws.append(totals)
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        for col_idx in range(1, 10):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = thin_border
            if col_idx >= 4:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="center")

        # Auto column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_material_excel(report: MaterialReport) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Moddiy_Hisobot"

        title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
        sub_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid") # Dark green
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        ws.append([f"Moddiy Hisobot - {report.organization_name}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
        ws.append([])

        h1 = ["Mahsulot nomi", "Birlik", "Boshlang'ich qoldiq", "", "Kirim", "", "Chiqim", "", "O'rtacha narx", "Oxirgi qoldiq", ""]
        h2 = ["", "", "Miqdor", "Summa", "Miqdor", "Summa", "Miqdor", "Summa", "", "Miqdor", "Summa"]
        ws.append(h1)
        ws.append(h2)

        start_row = 4
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row+1, end_column=1)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row+1, end_column=2)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row, end_column=4)
        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row, end_column=6)
        ws.merge_cells(start_row=start_row, start_column=7, end_row=start_row, end_column=8)
        ws.merge_cells(start_row=start_row, start_column=9, end_row=start_row+1, end_column=9)
        ws.merge_cells(start_row=start_row, start_column=10, end_row=start_row, end_column=11)

        for r in range(start_row, start_row + 2):
            for c in range(1, 12):
                cell = ws.cell(row=r, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        curr_row = start_row + 2
        for it in report.items:
            row_data = [
                it.item_name,
                it.unit,
                float(it.initial_qty),
                float(it.initial_sum),
                float(it.inflow_qty),
                float(it.inflow_sum),
                float(it.outflow_qty),
                float(it.outflow_sum),
                float(it.avg_price),
                float(it.final_qty),
                float(it.final_sum),
            ]
            ws.append(row_data)
            for c in range(1, 12):
                cell = ws.cell(row=curr_row, column=c)
                cell.font = regular_font
                cell.border = thin_border
                if c >= 3:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right")
            curr_row += 1

        # Totals
        totals = [
            "JAMI", "",
            "", float(report.total_initial_sum),
            "", float(report.total_inflow_sum),
            "", float(report.total_outflow_sum),
            "",
            "", float(report.total_final_sum)
        ]
        ws.append(totals)
        for c in range(1, 12):
            cell = ws.cell(row=curr_row, column=c)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = thin_border
            if c in [4, 6, 8, 11]:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 2nd Sheet: MXIK bo'yicha Jamlama
        if getattr(report, "mxik_groups", None):
            ws_mxik = wb.create_sheet(title="MXIK_Jamlama")
            ws_mxik.views.sheetView[0].showGridLines = True
            ws_mxik.append([f"MXIK (IKPU) bo'yicha Tovar va Moddiy Qoldiqlar Jamlamasi - {report.organization_name}"])
            ws_mxik.cell(row=1, column=1).font = title_font
            ws_mxik.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
            ws_mxik.cell(row=2, column=1).font = sub_font
            ws_mxik.append([])

            # Hierarchical Headers: Row 4 and Row 5
            mxik_h1 = [
                "№", "MXIK (IKPU) Kodi", "Mahsulotlar toifasi / Namunasi", "Birligi",
                "Tovarlar turlari soni", "Boshlang'ich qoldiq", "", "Davr kirimi", "",
                "Davr chiqimi", "", "Yakuniy qoldiq", ""
            ]
            mxik_h2 = [
                "", "", "", "", "",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)"
            ]
            ws_mxik.append(mxik_h1)
            ws_mxik.append(mxik_h2)

            # Merges
            ws_mxik.merge_cells("A4:A5")
            ws_mxik.merge_cells("B4:B5")
            ws_mxik.merge_cells("C4:C5")
            ws_mxik.merge_cells("D4:D5")
            ws_mxik.merge_cells("E4:E5")
            ws_mxik.merge_cells("F4:G4")
            ws_mxik.merge_cells("H4:I4")
            ws_mxik.merge_cells("J4:K4")
            ws_mxik.merge_cells("L4:M4")

            for r in range(4, 6):
                for c in range(1, 14):
                    cell = ws_mxik.cell(row=r, column=c)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border

            m_row = 6
            for idx, g in enumerate(report.mxik_groups, 1):
                ws_mxik.append([
                    idx,
                    g.ikpu_code,
                    g.ikpu_name or "Nomsiz toifa",
                    g.unit or "dona",
                    g.items_count,
                    float(g.initial_qty),
                    float(g.initial_sum),
                    float(g.inflow_qty),
                    float(g.inflow_sum),
                    float(g.outflow_qty),
                    float(g.outflow_sum),
                    float(g.final_qty),
                    float(g.final_sum),
                ])
                for col_idx in range(1, 14):
                    cell = ws_mxik.cell(row=m_row, column=col_idx)
                    cell.font = regular_font
                    cell.border = thin_border
                    if col_idx in [1, 2, 4]:
                        cell.alignment = Alignment(horizontal="center")
                    elif col_idx == 3:
                        cell.alignment = Alignment(horizontal="left")
                    elif col_idx == 5:
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="center")
                    elif col_idx in [6, 8, 10, 12]:
                        # Tovarlar soni (miqdor)
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        # Summalar
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right")
                m_row += 1

            # MXIK Totals
            tot_items_count = sum(g.items_count for g in report.mxik_groups)
            tot_init_qty = sum(float(g.initial_qty) for g in report.mxik_groups)
            tot_inf_qty = sum(float(g.inflow_qty) for g in report.mxik_groups)
            tot_outf_qty = sum(float(g.outflow_qty) for g in report.mxik_groups)
            tot_fin_qty = sum(float(g.final_qty) for g in report.mxik_groups)

            ws_mxik.append([
                "JAMI:", "", "", "",
                tot_items_count,
                tot_init_qty,
                float(report.total_initial_sum),
                tot_inf_qty,
                float(report.total_inflow_sum),
                tot_outf_qty,
                float(report.total_outflow_sum),
                tot_fin_qty,
                float(report.total_final_sum),
            ])
            ws_mxik.merge_cells(start_row=m_row, start_column=1, end_row=m_row, end_column=4)

            for col_idx in range(1, 14):
                cell = ws_mxik.cell(row=m_row, column=col_idx)
                cell.font = bold_font
                cell.fill = total_fill
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx == 5:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx >= 6:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right")

            # Column dimensions
            col_widths = {
                1: 6,    # №
                2: 22,   # MXIK Kodi
                3: 38,   # Mahsulot toifasi
                4: 12,   # Birligi
                5: 20,   # Tovarlar turlari soni
                6: 18,   # Boshlang'ich tovarlar soni
                7: 20,   # Boshlang'ich summa
                8: 18,   # Kirim tovarlar soni
                9: 20,   # Kirim summa
                10: 18,  # Chiqim tovarlar soni
                11: 20,  # Chiqim summa
                12: 18,  # Yakuniy tovarlar soni
                13: 22,  # Yakuniy summa
            }
            for col_i, w in col_widths.items():
                ws_mxik.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = w

            ws_mxik.freeze_panes = "A6"

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_pdf(title: str, headers: List[str], rows: List[List[Any]]) -> io.BytesIO:
        """Generic clean landscape PDF generator using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(name="DocTitle", parent=styles["Heading1"], fontSize=14, leading=18, textColor=colors.HexColor("#1E3A8A"))
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 12))

        table_data = [headers] + rows
        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(t)
        doc.build(elements)
        buffer.seek(0)
        return buffer
