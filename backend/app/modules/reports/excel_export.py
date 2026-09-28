import io
from decimal import Decimal
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.report import TrialBalanceReport, MaterialReport, AktSverkaReport

class ExcelExportEngine:
    """
    Styled .xlsx generator with:
    - Frozen header panes
    - Formatted currency cells (#,##0.00)
    - Bold total summary rows
    - Auto-fitted column widths
    """

    CURRENCY_FORMAT = "#,##0.00"
    QTY_FORMAT = "#,##0.000"

    @classmethod
    def _apply_styles(cls, ws, freeze_cell: str = "A6"):
        # Freeze panes
        ws.freeze_panes = freeze_cell

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged title cells for length calculation
                if cell.row < 4:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def export_trial_balance(cls, report: TrialBalanceReport) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Oborotka_OSV"

        # Styles
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        # Title
        ws.append([f"Aylanma Qoldiq Vedomosti (OSV) - {report.organization_name}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
        ws.cell(row=2, column=1).font = sub_font
        ws.append([])

        # Multi-level Headers
        headers_row1 = ["Schot", "Schot nomi", "Turi", "Boshlang'ich qoldiq", "", "Davr oboroti", "", "Oxirgi qoldiq", ""]
        headers_row2 = ["", "", "", "Debet", "Kredit", "Debet", "Kredit", "Debet", "Kredit"]
        ws.append(headers_row1)
        ws.append(headers_row2)

        start_row = 4
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row+1, end_column=1)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row+1, end_column=2)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row+1, end_column=3)
        ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=5)
        ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row, end_column=7)
        ws.merge_cells(start_row=start_row, start_column=8, end_row=start_row, end_column=9)

        for r in range(start_row, start_row + 2):
            for c in range(1, 10):
                cell = ws.cell(row=r, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

        # Rows
        for item in report.items:
            row_vals = [
                item.account_code,
                item.account_name,
                item.account_type,
                float(item.initial_debit),
                float(item.initial_credit),
                float(item.period_debit),
                float(item.period_credit),
                float(item.final_debit),
                float(item.final_credit)
            ]
            ws.append(row_vals)
            curr_row = ws.max_row
            for c_idx in range(1, 10):
                cell = ws.cell(row=curr_row, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
                if c_idx in [1, 3]:
                    cell.alignment = Alignment(horizontal="center")
                elif c_idx >= 4:
                    cell.number_format = cls.CURRENCY_FORMAT
                    cell.alignment = Alignment(horizontal="right")

        # Total Row
        tot_row_vals = [
            "JAMI",
            "Jami aylanma va qoldiqlar",
            "",
            float(report.total_initial_debit),
            float(report.total_initial_credit),
            float(report.total_period_debit),
            float(report.total_period_credit),
            float(report.total_final_debit),
            float(report.total_final_credit)
        ]
        ws.append(tot_row_vals)
        t_row = ws.max_row
        for c_idx in range(1, 10):
            cell = ws.cell(row=t_row, column=c_idx)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = thin_border
            if c_idx >= 4:
                cell.number_format = cls.CURRENCY_FORMAT
                cell.alignment = Alignment(horizontal="right")

        cls._apply_styles(ws, freeze_cell="A6")
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def export_material_report(cls, report: MaterialReport) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Moddiy_Hisobot"

        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        ws.append([f"Moddiy Hisobot (Material Stock Report) - {report.organization_name}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
        ws.cell(row=2, column=1).font = sub_font
        ws.append([])

        headers_row1 = ["№", "Tovar / Material nomi", "MXIK (IKPU)", "Birligi", "Boshlang'ich qoldiq", "", "Kirim", "", "Chiqim", "", "Oxirgi qoldiq", ""]
        headers_row2 = ["", "", "", "", "Miqdor", "Summa", "Miqdor", "Summa", "Miqdor", "Summa", "Miqdor", "Summa"]
        ws.append(headers_row1)
        ws.append(headers_row2)

        start_row = 4
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row+1, end_column=1)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row+1, end_column=2)
        ws.merge_cells(start_row=start_row, start_column=3, end_row=start_row+1, end_column=3)
        ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row+1, end_column=4)
        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row, end_column=6)
        ws.merge_cells(start_row=start_row, start_column=7, end_row=start_row, end_column=8)
        ws.merge_cells(start_row=start_row, start_column=9, end_row=start_row, end_column=10)
        ws.merge_cells(start_row=start_row, start_column=11, end_row=start_row, end_column=12)

        for r in range(start_row, start_row + 2):
            for c in range(1, 13):
                cell = ws.cell(row=r, column=c)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

        for idx, it in enumerate(report.items, 1):
            init_q = getattr(it, "initial_qty", getattr(it, "opening_qty", 0))
            init_s = getattr(it, "initial_sum", getattr(it, "opening_sum", 0))
            inf_q = getattr(it, "inflow_qty", 0)
            inf_s = getattr(it, "inflow_sum", 0)
            outf_q = getattr(it, "outflow_qty", 0)
            outf_s = getattr(it, "outflow_sum", 0)
            fin_q = getattr(it, "final_qty", getattr(it, "closing_qty", 0))
            fin_s = getattr(it, "final_sum", getattr(it, "closing_sum", 0))

            row_vals = [
                idx,
                it.item_name,
                it.ikpu_code or "",
                it.unit,
                float(init_q),
                float(init_s),
                float(inf_q),
                float(inf_s),
                float(outf_q),
                float(outf_s),
                float(fin_q),
                float(fin_s)
            ]
            ws.append(row_vals)
            curr_row = ws.max_row
            for c_idx in range(1, 13):
                cell = ws.cell(row=curr_row, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
                if c_idx in [1, 3, 4]:
                    cell.alignment = Alignment(horizontal="center")
                elif c_idx in [5, 7, 9, 11]:
                    cell.number_format = cls.QTY_FORMAT
                    cell.alignment = Alignment(horizontal="right")
                elif c_idx >= 6:
                    cell.number_format = cls.CURRENCY_FORMAT
                    cell.alignment = Alignment(horizontal="right")

        # Total Row
        tot_init = getattr(report, "total_initial_sum", getattr(report, "total_opening_sum", 0))
        tot_inf = getattr(report, "total_inflow_sum", 0)
        tot_outf = getattr(report, "total_outflow_sum", 0)
        tot_fin = getattr(report, "total_final_sum", getattr(report, "total_closing_sum", 0))

        tot_row = [
            "JAMI", "", "", "",
            "", float(tot_init),
            "", float(tot_inf),
            "", float(tot_outf),
            "", float(tot_fin)
        ]
        ws.append(tot_row)
        t_row = ws.max_row
        for c_idx in range(1, 13):
            cell = ws.cell(row=t_row, column=c_idx)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = thin_border
            if c_idx in [6, 8, 10, 12]:
                cell.number_format = cls.CURRENCY_FORMAT
                cell.alignment = Alignment(horizontal="right")

        cls._apply_styles(ws, freeze_cell="A6")

        # 2nd Sheet: MXIK bo'yicha Jamlama (Tovar soni va harakatlari)
        if getattr(report, "mxik_groups", None):
            ws_mxik = wb.create_sheet(title="MXIK_Jamlama")
            ws_mxik.views.sheetView[0].showGridLines = True
            ws_mxik.append([f"MXIK (IKPU) bo'yicha Tovar va Moddiy Qoldiqlar Jamlamasi - {report.organization_name}"])
            ws_mxik.cell(row=1, column=1).font = title_font
            ws_mxik.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
            ws_mxik.cell(row=2, column=1).font = sub_font
            ws_mxik.append([])

            # Hierarchical Headers: Row 4 and Row 5
            mxik_h1 = [
                "№", "MXIK (IKPU) Kodi", "Mahsulotlar toifasi / Namunasi", "Birligi",
                "Tovarlar turlari soni", "Boshlang'ich qoldiq", "", "Davr kirimi", "",
                "Davr chiqimi", "", "Yakuniy qoldiq", ""
            ]
            mxik_h2 = [
                "", "", "", "", "",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)",
                "Tovarlar soni", "Summasi (so'm)"
            ]
            ws_mxik.append(mxik_h1)
            ws_mxik.append(mxik_h2)

            ws_mxik.merge_cells("A4:A5")
            ws_mxik.merge_cells("B4:B5")
            ws_mxik.merge_cells("C4:C5")
            ws_mxik.merge_cells("D4:D5")
            ws_mxik.merge_cells("E4:E5")
            ws_mxik.merge_cells("F4:G4")
            ws_mxik.merge_cells("H4:I4")
            ws_mxik.merge_cells("J4:K4")
            ws_mxik.merge_cells("L4:M4")

            for r in range(4, 6):
                for c in range(1, 14):
                    cell = ws_mxik.cell(row=r, column=c)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = thin_border

            m_row = 6
            for idx, g in enumerate(report.mxik_groups, 1):
                ws_mxik.append([
                    idx,
                    g.ikpu_code,
                    g.ikpu_name or "Nomsiz toifa",
                    g.unit or "dona",
                    g.items_count,
                    float(g.initial_qty),
                    float(g.initial_sum),
                    float(g.inflow_qty),
                    float(g.inflow_sum),
                    float(g.outflow_qty),
                    float(g.outflow_sum),
                    float(g.final_qty),
                    float(g.final_sum),
                ])
                for col_idx in range(1, 14):
                    cell = ws_mxik.cell(row=m_row, column=col_idx)
                    cell.font = regular_font
                    cell.border = thin_border
                    if col_idx in [1, 2, 4]:
                        cell.alignment = Alignment(horizontal="center")
                    elif col_idx == 3:
                        cell.alignment = Alignment(horizontal="left")
                    elif col_idx == 5:
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="center")
                    elif col_idx in [6, 8, 10, 12]:
                        cell.number_format = cls.QTY_FORMAT
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.number_format = cls.CURRENCY_FORMAT
                        cell.alignment = Alignment(horizontal="right")
                m_row += 1

            # MXIK Totals
            tot_items_count = sum(g.items_count for g in report.mxik_groups)
            tot_init_qty = sum(float(g.initial_qty) for g in report.mxik_groups)
            tot_inf_qty = sum(float(g.inflow_qty) for g in report.mxik_groups)
            tot_outf_qty = sum(float(g.outflow_qty) for g in report.mxik_groups)
            tot_fin_qty = sum(float(g.final_qty) for g in report.mxik_groups)
            tot_init_sum = float(getattr(report, "total_initial_sum", getattr(report, "total_opening_sum", 0)))
            tot_inf_sum = float(report.total_inflow_sum)
            tot_outf_sum = float(report.total_outflow_sum)
            tot_fin_sum = float(getattr(report, "total_final_sum", getattr(report, "total_closing_sum", 0)))

            ws_mxik.append([
                "JAMI:", "", "", "",
                tot_items_count,
                tot_init_qty,
                tot_init_sum,
                tot_inf_qty,
                tot_inf_sum,
                tot_outf_qty,
                tot_outf_sum,
                tot_fin_qty,
                tot_fin_sum,
            ])
            ws_mxik.merge_cells(start_row=m_row, start_column=1, end_row=m_row, end_column=4)

            for col_idx in range(1, 14):
                cell = ws_mxik.cell(row=m_row, column=col_idx)
                cell.font = bold_font
                cell.fill = total_fill
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx == 5:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in [6, 8, 10, 12]:
                    cell.number_format = cls.QTY_FORMAT
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx >= 7:
                    cell.number_format = cls.CURRENCY_FORMAT
                    cell.alignment = Alignment(horizontal="right")

            cls._apply_styles(ws_mxik, freeze_cell="A6")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def export_akt_sverka(cls, report: AktSverkaReport) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Akt_Sverka"

        title_font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        regular_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        ws.append([f"Solishuv Dalolatnomasi (Akt Sverka)"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"{report.organization_name} va {report.counterparty_name} o'rtasida"])
        ws.cell(row=2, column=1).font = sub_font
        ws.append([f"Davr: {report.from_date.strftime('%d.%m.%Y')} dan {report.to_date.strftime('%d.%m.%Y')} gacha"])
        ws.cell(row=3, column=1).font = sub_font
        ws.append([])

        headers = ["Sana", "Hujjat raqami", "Turi", "Operatsiya mazmuni", "Debet (Bizning foydamizga)", "Kredit (Bizning qarzimiz)", "Joriy Saldo"]
        ws.append(headers)
        h_row = ws.max_row
        for c in range(1, 8):
            cell = ws.cell(row=h_row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Boshlang'ich qoldiq
        ws.append([
            report.from_date.strftime("%d.%m.%Y"),
            "-",
            "Boshlang'ich",
            "Davr boshiga qoldiq saldo",
            "",
            "",
            float(report.initial_balance)
        ])
        b_row = ws.max_row
        for c in range(1, 8):
            cell = ws.cell(row=b_row, column=c)
            cell.font = bold_font
            cell.border = thin_border
            if c == 7:
                cell.number_format = cls.CURRENCY_FORMAT
                cell.alignment = Alignment(horizontal="right")

        # Entries
        for it in report.items:
            ws.append([
                it.doc_date.strftime("%d.%m.%Y"),
                it.doc_number or "",
                it.doc_type,
                it.description or "",
                float(it.debit) if it.debit > 0 else "",
                float(it.credit) if it.credit > 0 else "",
                float(it.running_balance)
            ])
            curr_row = ws.max_row
            for c in range(1, 8):
                cell = ws.cell(row=curr_row, column=c)
                cell.font = regular_font
                cell.border = thin_border
                if c in [1, 2, 3]:
                    cell.alignment = Alignment(horizontal="center")
                elif c in [5, 6, 7]:
                    cell.number_format = cls.CURRENCY_FORMAT
                    cell.alignment = Alignment(horizontal="right")

        # Final totals
        ws.append([
            "JAMI",
            "",
            "",
            "Davr bo'yicha aylanma va yakuniy qoldiq",
            float(report.total_debit),
            float(report.total_credit),
            float(report.final_balance)
        ])
        f_row = ws.max_row
        for c in range(1, 8):
            cell = ws.cell(row=f_row, column=c)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = thin_border
            if c in [5, 6, 7]:
                cell.number_format = cls.CURRENCY_FORMAT
                cell.alignment = Alignment(horizontal="right")

        cls._apply_styles(ws, freeze_cell="A6")
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
