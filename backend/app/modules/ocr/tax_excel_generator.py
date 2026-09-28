import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from decimal import Decimal
from .tax_audit_schemas import TaxAuditDocument

class TaxAuditExcelGenerator:
    @staticmethod
    def generate_workbook(doc: TaxAuditDocument) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Киримсиз товарлар таҳлили"
        ws.views.sheetView[0].showGridLines = True

        NAVY = "1B365D"
        BLUE = "4A90E2"
        LIGHT_BG = "E8F1F5"
        ZEBRA = "F9FBFC"
        ALERT_RED = "FFF0F0"
        ALERT_BORDER = "FF4D4F"
        BORDER_CLR = "D3D3D3"

        font_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_sub = Font(name="Calibri", size=9, bold=True, color="1B365D")
        font_data = Font(name="Calibri", size=10)
        font_bold = Font(name="Calibri", size=10, bold=True)
        font_penalty_title = Font(name="Calibri", size=11, bold=True, color="1B365D")
        font_penalty_val = Font(name="Calibri", size=11, bold=True, color="B22222")

        thin = Border(left=Side(style='thin', color=BORDER_CLR),
                      right=Side(style='thin', color=BORDER_CLR),
                      top=Side(style='thin', color=BORDER_CLR),
                      bottom=Side(style='thin', color=BORDER_CLR))
        
        double_bottom = Border(left=Side(style='thin', color=BORDER_CLR),
                               right=Side(style='thin', color=BORDER_CLR),
                               top=Side(style='thin', color=BORDER_CLR),
                               bottom=Side(style='double', color=NAVY))

        # 1. Main Title
        ws.merge_cells("A1:M1")
        ws["A1"] = f'"{doc.company_name}" (СТИР: {doc.company_inn}) томонидан {doc.audit_year} йил давомида киримсиз сотилган товарлар таҳлили'
        ws["A1"].font = font_title
        ws["A1"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # 2. Subtitle
        ws.merge_cells("A2:M2")
        ws["A2"] = "Манба: Камерал солиқ текшируви ҳисоботидан автоматик рақамлаштирилган ва математик текширилган"
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="666666")
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        # 3. Hierarchical Headers
        headers = [
            ("A3:A4", "№", BLUE),
            ("B3:B4", "Товар номи", BLUE),
            ("C3:D3", f"{doc.audit_year} йил 1 январ қолдиқ", NAVY),
            ("E3:F3", f"{doc.audit_year} йил давомида кирим", NAVY),
            ("G3:I3", f"{doc.audit_year} йил сотилган (чиқим)", NAVY),
            ("J3:K3", f"{doc.audit_year} йил охирига қолдиқ", NAVY),
            ("L3:M3", "Фарқи (Киримсиз сотув)", "B22222"),
        ]
        for rng, txt, clr in headers:
            ws.merge_cells(rng)
            top_cell = rng.split(":")[0]
            ws[top_cell].value = txt
            ws[top_cell].font = font_header
            ws[top_cell].fill = PatternFill(start_color=clr, end_color=clr, fill_type="solid")
            ws[top_cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        subs = [
            ("C4", "Сони"), ("D4", "Суммаси"),
            ("E4", "Сони"), ("F4", "Суммаси"),
            ("G4", "Сони"), ("H4", "Ўртача нархи"), ("I4", "Суммаси"),
            ("J4", "Сони"), ("K4", "Суммаси"),
            ("L4", "Сони"), ("M4", "Киримсиз сумма")
        ]
        for col, txt in subs:
            ws[col].value = txt
            ws[col].font = font_sub
            ws[col].fill = PatternFill(start_color=LIGHT_BG, end_color=LIGHT_BG, fill_type="solid")
            ws[col].alignment = Alignment(horizontal="center", vertical="center")

        for r in range(3, 5):
            ws.row_dimensions[r].height = 24
            for c in range(1, 14):
                ws.cell(row=r, column=c).border = thin

        # 4. Data Injection
        start_row = 5
        items_count = len(doc.items)
        for idx, item in enumerate(doc.items, 1):
            curr_row = start_row + idx - 1
            ws.row_dimensions[curr_row].height = 20
            is_zebra = (idx % 2 == 0)
            z_fill = PatternFill(start_color=ZEBRA, end_color=ZEBRA, fill_type="solid") if is_zebra else None

            row_data = [
                (1, idx, "@", Alignment(horizontal="center")),
                (2, item.item_name, "@", Alignment(horizontal="left")),
                (3, float(item.opening_qty), "#,##0.00", Alignment(horizontal="right")),
                (4, float(item.opening_amount), "#,##0.00", Alignment(horizontal="right")),
                (5, float(item.inflow_qty), "#,##0.00", Alignment(horizontal="right")),
                (6, float(item.inflow_amount), "#,##0.00", Alignment(horizontal="right")),
                (7, float(item.sold_qty), "#,##0.00", Alignment(horizontal="right")),
                (8, float(item.avg_price), "#,##0.00", Alignment(horizontal="right")),
                (9, float(item.sold_amount), "#,##0.00", Alignment(horizontal="right")),
                (10, float(item.closing_qty), "#,##0.00", Alignment(horizontal="right")),
                (11, float(item.closing_amount), "#,##0.00", Alignment(horizontal="right")),
                (12, float(item.diff_qty), "#,##0.00", Alignment(horizontal="right")),
                (13, float(item.diff_amount), "#,##0.00", Alignment(horizontal="right")),
            ]

            for c_idx, val, fmt, align in row_data:
                c = ws.cell(row=curr_row, column=c_idx, value=val)
                c.font = font_data
                c.number_format = fmt
                c.alignment = align
                c.border = thin
                if c_idx in [12, 13] and val > 0:
                    c.fill = PatternFill(start_color=ALERT_RED, end_color=ALERT_RED, fill_type="solid")
                    c.font = font_bold
                elif z_fill:
                    c.fill = z_fill

        last_row = start_row + max(1, items_count) - 1
        tot_row = last_row + 1

        # 5. Grand Totals Row
        ws.merge_cells(f"A{tot_row}:B{tot_row}")
        ws[f"A{tot_row}"] = "ЖАМИ ҲИСОБОТ БЎЙИЧА:"
        ws[f"A{tot_row}"].font = font_header
        ws[f"A{tot_row}"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        ws[f"A{tot_row}"].alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[tot_row].height = 24

        for c_idx in range(1, 14):
            c = ws.cell(row=tot_row, column=c_idx)
            c.border = double_bottom
            c.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
            c.font = font_header
            if c_idx >= 3:
                col_letter = get_column_letter(c_idx)
                if c_idx == 8:
                    c.value = f"=AVERAGE({col_letter}{start_row}:{col_letter}{last_row})"
                else:
                    c.value = f"=SUM({col_letter}{start_row}:{col_letter}{last_row})"
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")

        # 6. Tax Penalty Summary Block
        if doc.summary:
            p_start = tot_row + 3
            ws.cell(row=p_start - 1, column=1, value="СОЛИҚ ВА ЖАРИМА ҲИСОБ-КИТОБИ (КАМЕРАЛ ТАҲЛИЛ НАТИЖАСИ):").font = font_penalty_title

            penalty_rows = [
                ("Жами киримсиз сотилган товарлар суммаси:", float(doc.summary.total_discrepancy_amount), "#,##0.00", False),
                ("ҚҚС (12%):", float(doc.summary.vat_amount), "#,##0.00", False),
                ("ҚҚСсиз солиқ солинадиган соф база:", float(doc.summary.net_tax_base), "#,##0.00", False),
                ("Фойда солиғига қўшимча (15%):", float(doc.summary.profit_tax_addition), "#,##0.00", False),
                ("Молиявий жарима (20%):", float(doc.summary.financial_penalty), "#,##0.00", False),
                ("ДАВЛАТ БЮДЖЕТИГА ЖАМИ ТЎЛОВ:", float(doc.summary.total_budget_liability), "#,##0.00", True),
            ]

            card_border = Border(
                left=Side(style='thin', color="B22222"),
                right=Side(style='thin', color="B22222"),
                top=Side(style='thin', color="B22222"),
                bottom=Side(style='thin', color="B22222")
            )

            for p_idx, (label, val, fmt, is_total) in enumerate(penalty_rows):
                r = p_start + p_idx
                ws.row_dimensions[r].height = 22
                ws.merge_cells(f"A{r}:E{r}")
                lbl_cell = ws.cell(row=r, column=1, value=label)
                lbl_cell.font = font_bold if is_total else font_data
                lbl_cell.alignment = Alignment(horizontal="left", vertical="center")

                val_cell = ws.cell(row=r, column=6, value=val)
                val_cell.font = font_penalty_val if is_total else font_bold
                val_cell.number_format = fmt
                val_cell.alignment = Alignment(horizontal="right", vertical="center")

                if is_total:
                    for c_i in range(1, 7):
                        cell = ws.cell(row=r, column=c_i)
                        cell.fill = PatternFill(start_color="FFF0F0", end_color="FFF0F0", fill_type="solid")
                        cell.border = card_border

        # 7. Column Widths
        widths = {1: 6, 2: 46, 3: 13, 4: 17, 5: 13, 6: 18, 7: 13, 8: 15, 9: 18, 10: 13, 11: 18, 12: 13, 13: 20}
        for col_i, w in widths.items():
            ws.column_dimensions[get_column_letter(col_i)].width = w

        out = BytesIO()
        wb.save(out)
        return out.getvalue()
