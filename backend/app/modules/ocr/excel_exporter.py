import io
from datetime import date, datetime
from decimal import Decimal
from typing import Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.modules.ocr.ocr_extractor import ExtractedDocument, ExtractedLineItem

class OCRExcelExporter:
    """
    Exports recovered/verified OCR invoice documents into an authentic,
    publication-grade Uzbekistan Electronic Invoice (Didox EHF / Hisobvaraq-faktura)
    in .xlsx format.
    
    Compatible with:
    1. Human accounting inspection and print/export ("aslidek qilib").
    2. Automatic ingestion back through DidoxParser / Documents pipeline ("keyin o'zim tekshirib kiritaman").
    """

    CURRENCY_FORMAT = "#,##0.00"
    QTY_FORMAT = "#,##0.000"
    DATE_FORMAT = "dd.mm.yyyy"

    @classmethod
    def export_document_to_excel(cls, doc: ExtractedDocument) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Hisobvaraq_Faktura"
        ws.views.sheetView[0].showGridLines = True

        # Palettes and fonts
        font_main = "Segoe UI"
        title_font = Font(name=font_main, size=13, bold=True, color="1E3A8A")
        subtitle_font = Font(name=font_main, size=9, bold=True, color="475569")
        meta_label_font = Font(name=font_main, size=9, bold=True, color="1E293B")
        meta_val_font = Font(name=font_main, size=9, bold=False, color="0F172A")
        header_font = Font(name=font_main, size=9, bold=True, color="FFFFFF")
        data_font = Font(name=font_main, size=9, bold=False, color="1E293B")
        total_font = Font(name=font_main, size=9, bold=True, color="0F172A")
        signature_font = Font(name=font_main, size=9, italic=True, color="334155")

        # Fills
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        card_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        total_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

        # Borders
        thin_side = Side(style="thin", color="CBD5E1")
        double_side = Side(style="double", color="1E3A8A")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        total_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=double_side)

        # Alignments
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # -------------------------------------------------------------
        # 1. Document Title & Subtitle Banner
        # -------------------------------------------------------------
        doc_no = doc.doc_number or "1"
        doc_date_val = doc.doc_date or date.today()
        doc_date_str = doc_date_val.strftime("%d.%m.%Y") if isinstance(doc_date_val, (date, datetime)) else str(doc_date_val)
        contract_no = doc.contract_number or "1"
        contract_date_str = doc.contract_date.strftime("%d.%m.%Y") if doc.contract_date else doc_date_str

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18

        ws.merge_cells("A1:P1")
        title_cell = ws["A1"]
        title_cell.value = f"ELEKTRON HISOBVARAQ-FAKTURA (EHF) № {doc_no}"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:P2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Sana: {doc_date_str} yil | Shartnoma raqami: {contract_no} ({contract_date_str}) | Turi: {doc.doc_type or 'EHF'}"
        sub_cell.font = subtitle_font
        sub_cell.alignment = Alignment(horizontal="left", vertical="center")

        # -------------------------------------------------------------
        # 2. Parties Info Cards (Supplier & Buyer Requisites)
        # -------------------------------------------------------------
        ws.row_dimensions[3].height = 8  # blank separator
        ws.row_dimensions[4].height = 18
        ws.row_dimensions[5].height = 18

        # Supplier (Cols A to G)
        ws.merge_cells("A4:G4")
        sup_hdr = ws["A4"]
        sup_hdr.value = f"YETKAZIB BERUVCHI: {doc.supplier_name or 'Noma\'lum korxona'}"
        sup_hdr.font = meta_label_font
        sup_hdr.fill = card_fill
        sup_hdr.alignment = align_left

        ws.merge_cells("A5:G5")
        sup_inn = ws["A5"]
        sup_inn.value = f"STIR (INN): {doc.supplier_inn or 'N/A'}"
        sup_inn.font = meta_val_font
        sup_inn.fill = card_fill
        sup_inn.alignment = align_left

        # Buyer (Cols H to P)
        ws.merge_cells("H4:P4")
        buy_hdr = ws["H4"]
        buy_hdr.value = f"XARIDOR: {doc.buyer_name or 'Noma\'lum xaridor'}"
        buy_hdr.font = meta_label_font
        buy_hdr.fill = card_fill
        buy_hdr.alignment = align_left

        ws.merge_cells("H5:P5")
        buy_inn = ws["H5"]
        buy_inn.value = f"STIR (INN): {doc.buyer_inn or 'N/A'}"
        buy_inn.font = meta_val_font
        buy_inn.fill = card_fill
        buy_inn.alignment = align_left

        ws.row_dimensions[6].height = 8  # blank separator

        # -------------------------------------------------------------
        # 3. Authentic Table Headers (Exact Didox Columns Matching)
        # -------------------------------------------------------------
        headers = [
            "№",
            "Hujjat raqami",
            "Sana",
            "Yetkazib beruvchi Nomi",
            "Yetkazib beruvchi STIR",
            "Xaridor Nomi",
            "Xaridor STIR",
            "Tovarlar (xizmatlar) nomi",
            "IKPU / MXIK kodi",
            "O'lchov birligi",
            "Miqdori",
            "Narxi",
            "Yetkazib berish qiymati",
            "QQS stavkasi",
            "QQS summasi",
            "Jami qiymat"
        ]

        table_header_row = 7
        ws.row_dimensions[table_header_row].height = 28
        for col_idx, header_title in enumerate(headers, start=1):
            cell = ws.cell(row=table_header_row, column=col_idx, value=header_title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_header
            cell.border = cell_border

        # -------------------------------------------------------------
        # 4. Data Rows (One per line item)
        # -------------------------------------------------------------
        current_row = table_header_row + 1
        total_qty = Decimal("0")
        total_delivery = Decimal("0")
        total_vat = Decimal("0")
        total_sum = Decimal("0")

        for idx, item in enumerate(doc.line_items, start=1):
            ws.row_dimensions[current_row].height = 22

            qty = Decimal(str(item.quantity or "1"))
            price = Decimal(str(item.price or "0"))
            vat_rate = Decimal(str(item.vat_rate or "12"))
            vat_amount = Decimal(str(item.vat_amount or "0"))
            item_total = Decimal(str(item.total_amount or "0"))

            delivery_val = qty * price
            if item_total == Decimal("0") and delivery_val > Decimal("0"):
                item_total = delivery_val + vat_amount

            total_qty += qty
            total_delivery += delivery_val
            total_vat += vat_amount
            total_sum += item_total

            row_values = [
                idx,                                          # 1. №
                doc_no,                                       # 2. Hujjat raqami
                doc_date_str,                                 # 3. Sana
                doc.supplier_name or "",                      # 4. Yetkazib beruvchi Nomi
                str(doc.supplier_inn or ""),                  # 5. Yetkazib beruvchi STIR
                doc.buyer_name or "",                         # 6. Xaridor Nomi
                str(doc.buyer_inn or ""),                     # 7. Xaridor STIR
                item.item_name,                               # 8. Tovarlar (xizmatlar) nomi
                str(item.ikpu_code or ""),                    # 9. IKPU / MXIK kodi
                item.unit or "dona",                          # 10. O'lchov birligi
                float(qty),                                   # 11. Miqdori
                float(price),                                 # 12. Narxi
                float(delivery_val),                          # 13. Yetkazib berish qiymati
                float(vat_rate),                              # 14. QQS stavkasi
                float(vat_amount),                            # 15. QQS summasi
                float(item_total)                             # 16. Jami qiymat
            ]

            for col_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = data_font
                cell.border = cell_border

                # Alignments and number formats
                if col_idx in [1, 3, 10]:
                    cell.alignment = align_center
                elif col_idx in [2, 4, 6, 8]:
                    cell.alignment = align_left
                elif col_idx in [5, 7, 9]:
                    cell.alignment = align_center
                    cell.data_type = "s"
                    cell.number_format = "@"  # Text format for INN / IKPU
                elif col_idx == 11:
                    cell.alignment = align_right
                    cell.number_format = cls.QTY_FORMAT
                elif col_idx in [12, 13, 15, 16]:
                    cell.alignment = align_right
                    cell.number_format = cls.CURRENCY_FORMAT
                elif col_idx == 14:
                    cell.alignment = align_right
                    cell.number_format = "0"

            current_row += 1

        # -------------------------------------------------------------
        # 5. Total Summary Row (Jami / Итого)
        # -------------------------------------------------------------
        ws.row_dimensions[current_row].height = 24
        
        # Merge columns A to J for "JAMI / ITOGO:" label
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        tot_label = ws.cell(row=current_row, column=1, value="JAMI (ITOGO):")
        tot_label.font = total_font
        tot_label.alignment = Alignment(horizontal="right", vertical="center")
        tot_label.fill = total_fill

        for c in range(1, 11):
            ws.cell(row=current_row, column=c).border = total_border
            ws.cell(row=current_row, column=c).fill = total_fill

        # Total Qty
        c_qty = ws.cell(row=current_row, column=11, value=float(total_qty))
        c_qty.font = total_font
        c_qty.fill = total_fill
        c_qty.alignment = align_right
        c_qty.number_format = cls.QTY_FORMAT
        c_qty.border = total_border

        # Empty Price cell
        c_empty = ws.cell(row=current_row, column=12, value="")
        c_empty.fill = total_fill
        c_empty.border = total_border

        # Total Delivery Value
        c_del = ws.cell(row=current_row, column=13, value=float(total_delivery))
        c_del.font = total_font
        c_del.fill = total_fill
        c_del.alignment = align_right
        c_del.number_format = cls.CURRENCY_FORMAT
        c_del.border = total_border

        # Empty VAT Rate cell
        c_vrate = ws.cell(row=current_row, column=14, value="")
        c_vrate.fill = total_fill
        c_vrate.border = total_border

        # Total VAT
        c_vat = ws.cell(row=current_row, column=15, value=float(total_vat))
        c_vat.font = total_font
        c_vat.fill = total_fill
        c_vat.alignment = align_right
        c_vat.number_format = cls.CURRENCY_FORMAT
        c_vat.border = total_border

        # Total Grand Total
        c_tot = ws.cell(row=current_row, column=16, value=float(total_sum))
        c_tot.font = Font(name=font_main, size=10, bold=True, color="1E3A8A")
        c_tot.fill = total_fill
        c_tot.alignment = align_right
        c_tot.number_format = cls.CURRENCY_FORMAT
        c_tot.border = total_border

        current_row += 2

        # -------------------------------------------------------------
        # 6. Authentic Signature Block (Rahbar & Bosh buxgalter)
        # -------------------------------------------------------------
        ws.row_dimensions[current_row].height = 20
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        sig_left = ws.cell(row=current_row, column=1, value="Rahbar (Yetkazib beruvchi): ____________________ (imzo)   M.O'.")
        sig_left.font = signature_font
        sig_left.alignment = align_left

        ws.merge_cells(start_row=current_row, start_column=8, end_row=current_row, end_column=16)
        sig_right = ws.cell(row=current_row, column=8, value="Bosh buxgalter: ____________________ (imzo)")
        sig_right.font = signature_font
        sig_right.alignment = align_left

        current_row += 1
        ws.row_dimensions[current_row].height = 20
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        rec_left = ws.cell(row=current_row, column=1, value="Qabul qilib oluvchi (Xaridor): ____________________ (imzo)")
        rec_left.font = signature_font
        rec_left.alignment = align_left

        # -------------------------------------------------------------
        # 7. Column Widths & Panes
        # -------------------------------------------------------------
        ws.freeze_panes = f"A{table_header_row + 1}"

        min_widths = {
            1: 5,   # №
            2: 14,  # Hujjat raqami
            3: 12,  # Sana
            4: 26,  # Yetkazib beruvchi Nomi
            5: 16,  # Yetkazib beruvchi STIR
            6: 26,  # Xaridor Nomi
            7: 16,  # Xaridor STIR
            8: 34,  # Tovarlar nomi
            9: 20,  # IKPU
            10: 10, # Birlik
            11: 12, # Miqdori
            12: 14, # Narxi
            13: 16, # Yetkazib berish qiymati
            14: 12, # QQS stavkasi
            15: 14, # QQS summasi
            16: 18  # Jami qiymat
        }

        for col_idx, col_w in min_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = col_w

        # Save to buffer
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
